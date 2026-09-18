#!/usr/bin/env python3
"""Stage 1: extract page-level text from input/ into extracted/{type}/*.json.

Electronic files are read from their text layer. Pages without a text layer
(scanned pages) are OCR-ed when Tesseract is installed; otherwise they are
recorded under manifest "needs_ocr" so the run stays auditable and rerunnable.
"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
INPUT_DIR = BASE / "input"
EXTRACTED_DIR = BASE / "extracted"

MIN_TEXT_CHARS = 20  # fewer characters than this on a PDF page means likely scanned
MAX_ENTRY_CHARS = 5000

DOC_TYPE_KEYWORDS = (
    ("drawings", ("drawing", "drw", "圖")),
    ("boq", ("boq", "報價", "工料", "數量")),
    ("specifications", ("spec", "規範")),
    ("meeting", ("meeting", "minute", "會議")),
    ("contract", ("contract", "agreement", "sow", "合約", "契約")),
)

DOC_TYPE_PREFIX = {
    "contract": "CON",
    "drawings": "DRW",
    "boq": "BOQ",
    "specifications": "SPEC",
    "meeting": "MTG",
    "general": "GEN",
}

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".txt", ".md"}


def detect_doc_type(filename: str) -> str:
    lowered = filename.lower()
    for doc_type, keywords in DOC_TYPE_KEYWORDS:
        if any(keyword in lowered for keyword in keywords):
            return doc_type
    return "general"


def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def ocr_pdf_page(page: "pymupdf.Page") -> str:
    import io

    import pytesseract
    from PIL import Image

    pix = page.get_pixmap(dpi=200)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    return pytesseract.image_to_string(image, lang="chi_tra+eng").strip()


def extract_pdf(path: Path) -> list[dict[str, object]]:
    import pymupdf

    entries: list[dict[str, object]] = []
    with pymupdf.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < MIN_TEXT_CHARS:
                if ocr_available():
                    try:
                        text = ocr_pdf_page(page)
                    except Exception as exc:  # OCR failure should not stop the run
                        entries.append(
                            {
                                "location": f"第 {index} 頁",
                                "text": "",
                                "needs_ocr": True,
                                "note": f"OCR failed: {exc}",
                            }
                        )
                        continue
                else:
                    entries.append(
                        {"location": f"第 {index} 頁", "text": "", "needs_ocr": True}
                    )
                    continue
            entries.append({"location": f"第 {index} 頁", "text": text})
    return entries


def extract_docx(path: Path) -> list[dict[str, object]]:
    import docx

    document = docx.Document(str(path))
    entries: list[dict[str, object]] = []
    paragraph_no = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_no += 1
        entries.append({"location": f"段落 {paragraph_no}", "text": text})
    for table_no, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(row for row in rows if row.strip(" |"))
        if text:
            entries.append({"location": f"表格 {table_no}", "text": text})
    return entries


def extract_xlsx(path: Path) -> list[dict[str, object]]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    entries: list[dict[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    lines.append(" | ".join(cells))
                if sum(len(line) for line in lines) > MAX_ENTRY_CHARS:
                    break
            text = "\n".join(lines)
            if text:
                entries.append({"location": f"工作表 {sheet.title}", "text": text})
    finally:
        workbook.close()
    return entries


def extract_text_file(path: Path) -> list[dict[str, object]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    entries: list[dict[str, object]] = []
    block_no = 0
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        block_no += 1
        entries.append({"location": f"區塊 {block_no}", "text": block})
    return entries


def extract_file(path: Path) -> list[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix in {".txt", ".md"}:
        return extract_text_file(path)
    raise ValueError(f"unsupported file type: {suffix}")


def clean_previous_output() -> None:
    if EXTRACTED_DIR.exists():
        for path in EXTRACTED_DIR.rglob("*.json"):
            path.unlink()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", default=str(INPUT_DIR), help="Directory of source documents")
    args = parser.parse_args()

    input_root = Path(args.input_dir).resolve()
    if not input_root.exists():
        raise SystemExit(f"input directory not found: {input_root}")

    clean_previous_output()
    EXTRACTED_DIR.mkdir(exist_ok=True)

    index: list[dict[str, object]] = []
    failed_files: list[dict[str, str]] = []
    needs_ocr: list[dict[str, str]] = []
    ocr_enabled = ocr_available()

    files = sorted(
        path for path in input_root.rglob("*") if path.is_file() and path.name != ".gitkeep"
    )
    for path in files:
        relative = path.relative_to(input_root).as_posix()
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            failed_files.append({"file": relative, "reason": "unsupported file type"})
            continue
        doc_type = detect_doc_type(path.name)
        try:
            entries = extract_file(path)
        except Exception as exc:
            failed_files.append({"file": relative, "reason": str(exc)})
            continue

        for entry in entries:
            if entry.get("needs_ocr"):
                needs_ocr.append({"file": relative, "location": str(entry["location"])})
            if len(str(entry["text"])) > MAX_ENTRY_CHARS:
                entry["text"] = str(entry["text"])[:MAX_ENTRY_CHARS] + "…（截斷）"

        record_name = relative.removesuffix(path.suffix).replace("/", "__")
        out_dir = EXTRACTED_DIR / doc_type
        out_dir.mkdir(parents=True, exist_ok=True)
        record = {
            "source_file": relative,
            "doc_type": doc_type,
            "finding_prefix": DOC_TYPE_PREFIX[doc_type],
            "ocr_used": ocr_enabled,
            "entries": entries,
        }
        out_path = out_dir / f"{record_name}.json"
        out_path.write_text(
            json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        index.append(
            {
                "source_file": relative,
                "doc_type": doc_type,
                "record": out_path.relative_to(BASE).as_posix(),
                "entries": len(entries),
                "needs_ocr_pages": sum(1 for entry in entries if entry.get("needs_ocr")),
            }
        )

    (EXTRACTED_DIR / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    manifest_path = BASE / "manifest.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["processed_files"] = len(index)
        manifest["failed_files"] = [item["file"] for item in failed_files]
        manifest["needs_ocr"] = needs_ocr
        manifest["extracted_index"] = "extracted/index.json"
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

    print(
        f"Extract: {len(index)} file(s) processed, {len(failed_files)} failed, "
        f"{len(needs_ocr)} page(s) waiting for OCR"
    )
    if needs_ocr and not ocr_enabled:
        print(
            "OCR is not configured. Install Tesseract (Windows: UB-Mannheim build) with "
            "chi_tra language data to extract scanned pages."
        )
    for item in failed_files:
        print(f"  FAILED {item['file']}: {item['reason']}")


if __name__ == "__main__":
    main()
