from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


PROJECT = Path(__file__).resolve().parents[1]
RUNTIME_FILES = ("manifest.json", "findings.json", "unresolved.json")


class PipelineTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.project = Path(self.temp_dir.name) / "project"
        shutil.copytree(PROJECT, self.project)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def run_script(
        self, name: str, *args: str, env: dict[str, str] | None = None
    ) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(self.project / "scripts" / name), *args],
            cwd=self.project,
            text=True,
            capture_output=True,
            check=False,
            env=env,
        )

    def clear_runtime_state(self) -> None:
        """Remove local artifacts so pipeline tests stay deterministic."""
        for name in ("input", "extracted", "normalized", "evidence", "reports"):
            target = self.project / name
            if not target.exists():
                continue
            for path in target.iterdir():
                if path.is_file():
                    path.unlink()
                elif path.is_dir():
                    shutil.rmtree(path)
        findings_dir = self.project / "findings"
        findings_dir.mkdir(exist_ok=True)
        for path in findings_dir.glob("*.json"):
            if path.name != "FINDING_EXAMPLE.json":
                path.unlink()
        for name in RUNTIME_FILES:
            target = self.project / name
            if target.exists():
                target.unlink()

    def write_input(self, name: str, content: str) -> Path:
        input_dir = self.project / "input"
        input_dir.mkdir(exist_ok=True)
        path = input_dir / name
        path.write_text(content, encoding="utf-8")
        return path

    def read_all_findings(self) -> list[dict[str, object]]:
        findings: list[dict[str, object]] = []
        for path in sorted((self.project / "findings").glob("*.json")):
            if path.name == "FINDING_EXAMPLE.json":
                continue
            findings.append(json.loads(path.read_text(encoding="utf-8")))
        return findings

    def write_finding(self, filename: str, **overrides: object) -> None:
        finding: dict[str, object] = {
            "finding_id": "CON-001",
            "classification": "責任限制",
            "judgment": "高風險",
            "risk_level": "high",
            "source_file": "contract.pdf",
            "location": "第 1 頁",
            "quote": "責任不受限制。",
            "analysis": "建議加入責任上限。",
            "status": "需人工確認",
            "config_rule": "risk_rules.yaml:high:賠償無上限",
        }
        finding.update(overrides)
        (self.project / "findings" / filename).write_text(
            json.dumps(finding, ensure_ascii=False), encoding="utf-8"
        )

    def test_evidence_backed_finding_reaches_report(self) -> None:
        self.write_finding("CON-001.json")
        self.assertEqual(self.run_script("00_manifest.py", "--project", "test-project").returncode, 0)
        self.assertEqual(self.run_script("04_evidence_check.py").returncode, 0)
        self.assertEqual(self.run_script("05_build_report.py").returncode, 0)

        findings = json.loads((self.project / "findings.json").read_text(encoding="utf-8"))
        unresolved = json.loads((self.project / "unresolved.json").read_text(encoding="utf-8"))
        report = (self.project / "reports" / "report.md").read_text(encoding="utf-8")
        self.assertEqual([item["finding_id"] for item in findings], ["CON-001"])
        self.assertEqual([item["finding_id"] for item in unresolved], ["CON-001"])
        self.assertIn("CON-001", report)

    def test_missing_evidence_is_blocked(self) -> None:
        self.write_finding("CON-002.json", finding_id="CON-002", quote="")
        result = self.run_script("04_evidence_check.py")
        self.assertEqual(result.returncode, 1)

        index = json.loads((self.project / "evidence" / "index.json").read_text(encoding="utf-8"))
        self.assertFalse(index[0]["has_evidence"])
        self.assertIn("quote", index[0]["missing"])

    def test_full_text_pipeline_reaches_report(self) -> None:
        self.clear_runtime_state()
        self.write_input(
            "sample_contract.txt",
            "第 1 條：乙方之賠償責任不受任何限制，且包含間接損害。\n\n"
            "第 2 條：本合約期間屆滿後自動續約，甲方得不另行通知。\n\n"
            "第 3 條：付款週期為請款後 90 日。\n",
        )
        for script in ("00_manifest.py", "01_extract.py", "02_normalize.py", "03_analyze.py"):
            self.assertEqual(
                self.run_script(script).returncode, 0, msg=f"{script} failed"
            )
        self.assertEqual(self.run_script("04_evidence_check.py").returncode, 0)
        self.assertEqual(self.run_script("05_build_report.py").returncode, 0)

        findings = json.loads((self.project / "findings.json").read_text(encoding="utf-8"))
        self.assertTrue(
            any("賠償責任不受任何限制" in str(item["quote"]) for item in findings),
            msg=f"expected verbatim quote in findings: {findings}",
        )
        self.assertTrue(all(item["source_file"] == "sample_contract.txt" for item in findings))
        self.assertEqual(
            [item["finding_id"] for item in findings],
            ["CON-001", "CON-002", "CON-003"],
        )

        unresolved = json.loads((self.project / "unresolved.json").read_text(encoding="utf-8"))
        self.assertEqual(len(unresolved), 3)

        manifest = json.loads((self.project / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["processed_files"], 1)
        self.assertEqual(manifest["failed_files"], [])
        self.assertIn("CON-001", (self.project / "reports" / "report.md").read_text(encoding="utf-8"))

    def test_docx_and_xlsx_extraction(self) -> None:
        self.clear_runtime_state()

        import docx
        import openpyxl

        document = docx.Document()
        document.add_paragraph("乙方之賠償責任不受任何限制，且包含間接損害。")
        document.save(str(self.project / "input" / "sample_contract.docx"))

        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "本合約期間屆滿後自動續約"
        workbook.save(str(self.project / "input" / "sample_boq.xlsx"))

        for script in ("00_manifest.py", "01_extract.py", "02_normalize.py", "03_analyze.py"):
            self.assertEqual(self.run_script(script).returncode, 0, msg=f"{script} failed")

        index = json.loads((self.project / "extracted" / "index.json").read_text(encoding="utf-8"))
        self.assertEqual(
            sorted(str(item["source_file"]) for item in index),
            ["sample_boq.xlsx", "sample_contract.docx"],
        )
        findings = self.read_all_findings()
        prefixes = {str(item["finding_id"]).split("-")[0] for item in findings}
        self.assertIn("CON", prefixes)
        self.assertIn("BOQ", prefixes)

    def test_scanned_pdf_marks_needs_ocr(self) -> None:
        self.clear_runtime_state()

        import pymupdf

        input_dir = self.project / "input"
        input_dir.mkdir(exist_ok=True)
        pdf_path = input_dir / "scanned_contract.pdf"
        doc = pymupdf.open()
        doc.new_page()
        doc.save(str(pdf_path))
        doc.close()

        self.assertEqual(self.run_script("00_manifest.py").returncode, 0)
        self.assertEqual(self.run_script("01_extract.py").returncode, 0)

        manifest = json.loads((self.project / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["processed_files"], 1)
        self.assertEqual(
            [item["file"] for item in manifest["needs_ocr"]], ["scanned_contract.pdf"]
        )
        extracted = json.loads((self.project / "extracted" / "index.json").read_text(encoding="utf-8"))
        self.assertEqual(extracted[0]["needs_ocr_pages"], 1)

    def test_mock_llm_proposes_extra_finding(self) -> None:
        self.clear_runtime_state()
        self.write_input(
            "sample_contract.txt",
            "第 1 條：乙方之賠償責任不受任何限制，且包含間接損害。\n",
        )
        for script in ("00_manifest.py", "01_extract.py", "02_normalize.py"):
            self.assertEqual(self.run_script(script).returncode, 0, msg=f"{script} failed")

        env = {**os.environ, "ZAI_MOCK": "1"}
        self.assertEqual(self.run_script("03_analyze.py", "--llm", env=env).returncode, 0)

        findings = self.read_all_findings()
        self.assertEqual(len(findings), 2)
        mock_findings = [item for item in findings if "MOCK" in str(item["analysis"])]
        self.assertEqual(len(mock_findings), 1)
        self.assertEqual(mock_findings[0]["config_rule"], "risk_rules.yaml:high:賠償無上限")

    def test_mock_llm_synthesis_writes_report(self) -> None:
        self.write_finding("CON-001.json")
        self.assertEqual(self.run_script("00_manifest.py", "--project", "test-project").returncode, 0)
        self.assertEqual(self.run_script("04_evidence_check.py").returncode, 0)
        self.assertEqual(self.run_script("05_build_report.py").returncode, 0)

        env = {**os.environ, "ZAI_MOCK": "1"}
        self.assertEqual(self.run_script("06_llm_synthesis.py", env=env).returncode, 0)
        synthesis = (self.project / "reports" / "synthesis.md").read_text(encoding="utf-8")
        self.assertIn("CON-001", synthesis)

    def test_llm_synthesis_requires_api_key(self) -> None:
        self.write_finding("CON-001.json")
        self.run_script("04_evidence_check.py")
        self.run_script("05_build_report.py")

        env = {key: value for key, value in os.environ.items() if key not in {"ZAI_API_KEY", "ZAI_MOCK"}}
        result = self.run_script("06_llm_synthesis.py", env=env)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("ZAI_API_KEY", result.stderr)


if __name__ == "__main__":
    unittest.main()
